# Copyright 2024-2026 OpenJC
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
#
# 本文件来源于 JC RAG Community Edition
# 开源地址: https://gitee.com/Open-JC/jc-rag | https://github.com/MetaJC/JC-RAG
# 二次开发请保留此出处声明。


"""PowerPoint PPTX 文件加载器"""
from __future__ import annotations

from typing import Any

from .base_loader import DocumentLoader


class PptxLoader(DocumentLoader):
    """PPTX 文件加载器"""
    
    def can_load(self, file_path: str) -> bool:
        """判断是否为 PPTX 文件"""
        return file_path.lower().endswith('.pptx')
    
    def load(self, file_path: str) -> str:
        """加载 PPTX 文件内容"""
        return self.load_with_meta(file_path)[0]
    
    def load_with_meta(self, file_path: str) -> tuple[str, dict[str, Any]]:
        """加载 PPTX 文件内容及元数据"""
        try:
            from pptx import Presentation
        except ImportError as ex:
            raise ValueError(
                "缺少 PPTX 解析依赖 python-pptx，请先执行 pip install python-pptx"
            ) from ex
        
        prs = Presentation(file_path)
        slides = []
        
        for slide_idx, slide in enumerate(prs.slides, 1):
            texts = []
            for shape in slide.shapes:
                if hasattr(shape, 'text'):
                    texts.append(shape.text)
            
            if texts:
                slides.append(f"=== 幻灯片 {slide_idx} ===\n" + '\n'.join(texts))
        
        content = '\n\n'.join(slides)
        meta = {
            'file_type': 'pptx',
            'slide_count': len(prs.slides),
            'char_count': len(content),
        }
        
        return content, meta
